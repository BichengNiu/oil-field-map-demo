"""Complete, source-count-checked PortWatch downloads with dated provenance."""
from __future__ import annotations

from collections import defaultdict
import csv
from datetime import date, datetime, timedelta, timezone
import hashlib
import io
import json
import logging
import math
import os
from pathlib import Path
import tempfile
import uuid
import zipfile
import zlib

import portwatch as pw
import data_store
from csv_export import csv_bytes, safe_csv_value
from portwatch_records import PORT_METRICS, CHOKE_METRICS, EXPORT_SHIP_LABELS as TYPES

logger = logging.getLogger(__name__)
VERSION = 1
FIRST_DAY = date(2019, 1, 1)
PAGE_SIZE = 1000
BATCH_SIZE = 60
XLSX_ROW_LIMIT = 20_000
INLINE_EXPORT_ROW_LIMIT = 20_000
KINDS = {"ports": "港口", "chokepoints": "咽喉要道"}
BASE_FIELDS = ["date", "portid", "portname", "country", "ISO3", "region", "lat", "lon"]
CATALOG_FIELDS = ["node_kind", "portid", "node_name", "country", "region", "lat", "lon",
                  "statistics_available", "first_available_utc", "latest_available_utc",
                  "source_url", "coverage_note"]
COVERAGE_FIELDS = ["node_kind", "portid", "node_name", "status", "first_available_utc",
                   "latest_available_utc", "export_first_utc", "export_last_utc", "records",
                   "requested_days", "missing_days", "missing_intervals_utc", "null_metric_cells"]


class DownloadError(ValueError):
    pass


class CacheMiss(DownloadError):
    """A read-only page requested data that is not present in local cache."""


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _day(value) -> date:
    if isinstance(value, (int, float)):
        return datetime.fromtimestamp(value / 1000, timezone.utc).date()
    return date.fromisoformat(str(value))


def node(kind: str, raw: dict) -> dict:
    if kind not in KINDS:
        raise DownloadError("数据类型无效")
    pid = str(raw["portid"])
    supported = pw.valid_port_ids((pid,)) if kind == "ports" else pw.valid_chokepoint_ids((pid,))
    return {"node_kind": kind, "portid": pid,
            "node_name": raw.get("name_cn") or raw.get("name") or raw.get("portname") or pid,
            "country": raw.get("country"),
            "region": raw.get("region") or raw.get("name_cn"),
            "lat": raw.get("lat"), "lon": raw.get("lon"),
            "statistics_available": supported,
            "source_url": raw.get("inventory_source_url") or raw.get("source_url") or pw.SOURCE,
            "coverage_note": raw.get("coverage_note", "AIS推算日度活动；覆盖受信号与识别方法影响")}


def _ids_where(ids: tuple[str, ...], kind: str) -> str:
    if kind not in ("ports", "chokepoints"):
        raise DownloadError("未知的PortWatch数据类型")
    valid = pw.valid_port_ids(ids) if kind == "ports" else pw.valid_chokepoint_ids(ids)
    if not ids or not valid or len(set(ids)) != len(ids):
        raise DownloadError("节点编号为空、重复或无效")
    return "portid IN (" + ",".join(f"'{pid}'" for pid in ids) + ")"


def bounds(kind: str, ids: tuple[str, ...]) -> tuple[dict, list[dict]]:
    output, queries = {}, []
    for offset in range(0, len(ids), BATCH_SIZE):
        batch = ids[offset:offset + BATCH_SIZE]
        where = _ids_where(batch, kind)
        statistics = [
            {"statisticType": "min", "onStatisticField": "date", "outStatisticFieldName": "first_date"},
            {"statisticType": "max", "onStatisticField": "date", "outStatisticFieldName": "last_date"},
            {"statisticType": "count", "onStatisticField": "ObjectId", "outStatisticFieldName": "records"},
        ]
        response = pw.query(pw.endpoint_for(kind), where=where, returnGeometry="false",
                             outStatistics=json.dumps(statistics), groupByFieldsForStatistics="portid",
                             resultRecordCount=PAGE_SIZE)
        if response.get("exceededTransferLimit"):
            raise DownloadError("节点日期范围查询被截断")
        for item in response.get("features", []):
            row = item["attributes"]
            pid = row["portid"]
            if pid not in batch or pid in output:
                raise DownloadError("日期范围查询返回未请求或重复节点")
            first, last = _day(row["first_date"]), _day(row["last_date"])
            if first > last or last > datetime.now(timezone.utc).date() or int(row["records"]) <= 0:
                raise DownloadError("源站日期范围或记录数无效")
            output[pid] = {"first": first.isoformat(), "last": last.isoformat(),
                           "records": int(row["records"])}
        queries.append({"endpoint": pw.endpoint_for(kind), "where": where, "operation": "node_date_bounds",
                        "retrieved_at_utc": _now()})
    return output, queries


def _cache_path() -> Path:
    return data_store.database_path()


def _cache(key: str, payload: dict | None = None,
           max_age_seconds: int = 86400) -> dict | None:
    """Cache only a validated complete query; a broken cache never hides source errors."""
    path = _cache_path()
    try:
        with data_store.connect(path) as conn:
            if payload is not None:
                conn.execute("INSERT OR REPLACE INTO portwatch_cache VALUES (?, ?, ?)",
                             (key, datetime.now(timezone.utc).timestamp(),
                              zlib.compress(json.dumps(payload, ensure_ascii=False, allow_nan=False).encode())))
                conn.execute("DELETE FROM portwatch_cache WHERE saved < ?", (datetime.now(timezone.utc).timestamp() - 7 * 86400,))
                return None
            row = conn.execute("SELECT saved,payload FROM portwatch_cache WHERE key=?", (key,)).fetchone()
        if row and datetime.now(timezone.utc).timestamp() - row[0] < max_age_seconds:
            return json.loads(zlib.decompress(row[1]))
    except (OSError, data_store.duckdb.Error, ValueError, zlib.error) as exc:
        logger.warning("PortWatch history cache unavailable (%s); continuing without it",
                       type(exc).__name__)
        pass
    return None


def fetch_window(kind: str, ids: tuple[str, ...], first: date, last: date,
                 force: bool = False, progress=None,
                 cache_revision: int | str = 0,
                 cache_ttl_seconds: int = 86400,
                 allow_stale: bool = False,
                 cache_only: bool = False) -> tuple[list[dict], dict]:
    where = f"{_ids_where(ids, kind)} AND date >= DATE '{first}' AND date <= DATE '{last}'"
    endpoint = pw.endpoint_for(kind)
    metrics = PORT_METRICS if kind == "ports" else CHOKE_METRICS
    source_fields = ["ObjectId", "date", "portid", "portname"] + (["country", "ISO3"] if kind == "ports" else []) + metrics
    key = hashlib.sha256(json.dumps(
        [VERSION, endpoint, where, source_fields, cache_revision,
         cache_ttl_seconds]).encode()).hexdigest()
    cache_age = 7 * 86400 if allow_stale else cache_ttl_seconds
    cached = None if force else _cache(key, max_age_seconds=cache_age)
    if cached:
        try:
            cached_at = datetime.fromisoformat(
                str(cached["query"]["retrieved_at_utc"]).replace("Z", "+00:00")
            )
            if cached_at.tzinfo is None:
                cached_at = cached_at.replace(tzinfo=timezone.utc)
            newest_fact = data_store.portwatch_window_latest_fetch(kind, ids, first, last)
        except (KeyError, TypeError, ValueError):
            newest_fact = None
            cached_at = None
        if cached_at is not None and (newest_fact is None or newest_fact <= cached_at):
            return cached["rows"], {**cached["query"], "cache_hit": True}
        if cached_at is not None:
            try:
                updates = data_store.cached_portwatch_rows_since(
                    kind, ids, first, last, cached_at
                )
            except Exception as exc:
                logger.warning("PortWatch newer local facts unavailable (%s)", type(exc).__name__)
                updates = []
            if updates:
                merged = {
                    (str(row.get("portid")), str(row.get("date"))[:10]): row
                    for row in cached["rows"]
                }
                for row in updates:
                    merged[(str(row["portid"]), str(row["date"])[:10])] = row
                merged_rows = [merged[key] for key in sorted(merged)]
                return merged_rows, {
                    **cached["query"],
                    "cache_hit": True,
                    "cache_source": "verified history with newer DuckDB facts",
                    "updated_local_rows": len(updates),
                    "source_count": len(merged_rows),
                    "exported_query_rows": len(merged_rows),
                }

    if not force:
        cached_rows = data_store.cached_portwatch_window(
            kind, ids, first, last, max_age_seconds=max(cache_ttl_seconds, 86_400),
            with_metadata=True, allow_stale=allow_stale)
        if cached_rows is not None:
            local_rows, retrieved_at = cached_rows
            return local_rows, {
                "endpoint": endpoint,
                "where": where,
                "retrieved_at_utc": retrieved_at.isoformat(),
                "cache_hit": True,
                "cache_source": "DuckDB verified daily facts",
                "source_count": len(local_rows),
                "exported_query_rows": len(local_rows),
                "count_verified": True,
            }

    if cache_only:
        raise CacheMiss("本地PortWatch历史缓存缺失或不完整")

    def count():
        value = pw.query(endpoint, where=where, returnGeometry="false", returnCountOnly="true")
        if "count" not in value:
            raise DownloadError("源站未返回用于核对分页的总记录数")
        return int(value["count"])

    expected = count()
    rows, keys = [], set()
    offset = 0
    while offset < expected:
        page = pw.query(endpoint, where=where, returnGeometry="false",
                         outFields=",".join(source_fields), orderByFields="portid ASC,date ASC,ObjectId ASC",
                         resultOffset=offset, resultRecordCount=PAGE_SIZE)
        items = page.get("features", [])
        if not items:
            raise DownloadError(f"{KINDS[kind]} {first}—{last} 分页提前结束（{offset}/{expected}）")
        for item in items:
            row = dict(item["attributes"])
            day = _day(row.get("date"))
            pid = row.get("portid")
            record_key = (pid, day)
            if pid not in ids or not first <= day <= last or record_key in keys:
                raise DownloadError(f"返回重复或超出请求范围的记录：{pid}/{day}")
            for metric in metrics:
                value = row.get(metric)
                if value is not None and (not isinstance(value, (int, float)) or
                                          not math.isfinite(value) or value < 0):
                    raise DownloadError(f"{pid}/{day} 的{metric}不是有效非负数值")
            row["date"] = day.isoformat()
            row.pop("ObjectId", None)
            rows.append(row)
            keys.add(record_key)
        offset += len(items)
        if progress:
            progress(f"{KINDS[kind]} {first}—{last}：{offset:,}/{expected:,} 条")
    if len(rows) != expected or count() != expected:
        raise DownloadError(f"{KINDS[kind]} {first}—{last} 源记录数发生变化或分页不完整，请重新生成")
    data_store.save_portwatch(kind, rows, endpoint)
    query = {"endpoint": endpoint, "where": where, "out_fields": source_fields,
             "retrieved_at_utc": _now(), "cache_hit": False, "source_count": expected,
             "exported_query_rows": len(rows), "count_verified": True}
    _cache(key, {"rows": rows, "query": query})
    return rows, query


def _years(first: date, last: date):
    current = first
    while current <= last:
        end = min(last, date(current.year, 12, 31))
        yield current, end
        current = end + timedelta(days=1)


def _refresh_segments(first: date, last: date, force: bool):
    """Split history into a cached stable range and exactly seven recent days."""
    if force:
        yield first, last, True
        return
    recent_start = datetime.now(timezone.utc).date() - timedelta(days=6)
    stable_end = recent_start - timedelta(days=1)
    if first <= min(last, stable_end):
        yield first, min(last, stable_end), False
    if last >= recent_start:
        yield max(first, recent_start), last, True


def _missing_intervals(first: date, last: date, days: set[date]) -> str:
    intervals, beginning = [], None
    current = first
    while current <= last:
        if current not in days and beginning is None:
            beginning = current
        if current in days and beginning is not None:
            intervals.append(f"{beginning}/{current - timedelta(days=1)}")
            beginning = None
        current += timedelta(days=1)
    if beginning is not None:
        intervals.append(f"{beginning}/{last}")
    return ";".join(intervals)


def _missing_intervals_bitmap(first: date | None, last: date | None,
                              present: bytearray) -> str:
    if first is None or last is None:
        return ""
    intervals, beginning = [], None
    for offset, value in enumerate(present):
        day = first + timedelta(days=offset)
        if not value and beginning is None:
            beginning = day
        if value and beginning is not None:
            intervals.append(f"{beginning}/{day - timedelta(days=1)}")
            beginning = None
    if beginning is not None:
        intervals.append(f"{beginning}/{last}")
    return ";".join(intervals)


def _export_directory() -> Path:
    path = Path(tempfile.gettempdir()) / "oil-field-map-exports"
    path.mkdir(parents=True, exist_ok=True)
    cutoff = datetime.now(timezone.utc).timestamp() - 6 * 3600
    for artifact in path.glob("portwatch-*.zip"):
        try:
            if artifact.stat().st_mtime < cutoff:
                artifact.unlink()
        except OSError:
            pass
    return path


def _write_csv_member(archive, name: str, rows, fields: list[str]) -> None:
    with archive.open(name, "w") as raw:
        text = io.TextIOWrapper(raw, encoding="utf-8-sig", newline="")
        writer = csv.DictWriter(text, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: safe_csv_value(row.get(key)) for key in fields})
        text.flush()
        text.detach()


def _derive(rows: list[dict], kind: str) -> None:
    metrics = ["portcalls", "import", "export"] if kind == "ports" else ["n_total", "capacity"]
    by_node = defaultdict(dict)
    for row in rows:
        by_node[row["portid"]][_day(row["date"])] = row
    for row in rows:
        day, series = _day(row["date"]), by_node[row["portid"]]
        for window in (7, 30):
            reports = [series.get(day - timedelta(days=i)) for i in range(window)]
            available = sum(report is not None for report in reports)
            row[f"observed_days_{window}d"] = available
            for metric in metrics:
                values = [report.get(metric) if report else None for report in reports]
                row[f"avg_{metric}_{window}d"] = (sum(values) / window if all(
                    value is not None for value in values) else None)


def _collect_streaming(prepared: dict, mode: str, first: date | None, last: date | None,
                       derived: bool, force: bool, progress=None) -> dict:
    """Write large exports a verified query chunk at a time to a temporary ZIP."""
    directory = _export_directory()
    artifact_path = directory / f"portwatch-{uuid.uuid4().hex}.zip"
    metadata, coverage, catalog = [], [], []
    row_counts = {}

    with zipfile.ZipFile(artifact_path, "w", compression=zipfile.ZIP_DEFLATED,
                         compresslevel=6) as archive:
        for kind in KINDS:
            if kind not in prepared:
                continue
            selected, limits, bound_queries = prepared[kind]
            metadata.extend(bound_queries)
            periods = defaultdict(list)
            wanted = {}
            stats = {}
            tails = {}
            metrics = PORT_METRICS if kind == "ports" else CHOKE_METRICS
            selected_by_id = {str(node["portid"]): node for node in selected}

            for node_row in selected:
                pid = str(node_row["portid"])
                limit = limits.get(pid)
                catalog.append({
                    **node_row,
                    "first_available_utc": limit["first"] if limit else None,
                    "latest_available_utc": limit["last"] if limit else None,
                })
                if mode == "latest":
                    begin = end = _day(limit["last"]) if limit else None
                elif mode == "all":
                    begin, end = (_day(limit["first"]), _day(limit["last"])) if limit else (None, None)
                else:
                    begin, end = first, last
                wanted[pid] = (begin, end)
                expected = (end - begin).days + 1 if begin and end else 0
                stats[pid] = {
                    "begin": begin, "end": end, "expected": expected,
                    "present": bytearray(expected), "records": 0,
                    "null_metric_cells": 0, "first": None, "last": None,
                }
                if not limit:
                    continue
                fetch_first = max(
                    _day(limit["first"]), begin - timedelta(days=29) if derived else begin
                )
                fetch_last = min(_day(limit["last"]), end)
                if fetch_first > fetch_last:
                    continue
                for segment_first, segment_last in _years(fetch_first, fetch_last):
                    for part_first, part_last, part_force in _refresh_segments(
                        segment_first, segment_last, force or mode == "latest"
                    ):
                        periods[(part_first, part_last, part_force)].append(pid)

            wrote_dataset = False
            chunk_number = 0
            for (segment_first, segment_last, period_force), pids in sorted(periods.items()):
                for offset in range(0, len(pids), BATCH_SIZE):
                    chunk = tuple(sorted(pids[offset:offset + BATCH_SIZE]))
                    rows, query = fetch_window(
                        kind, chunk, segment_first, segment_last,
                        force=period_force, progress=progress,
                    )
                    metadata.append(query)

                    derivation_rows = rows
                    if derived:
                        context_first = segment_first - timedelta(days=29)
                        context = [
                            row for pid in chunk for row in tails.get(pid, [])
                            if context_first <= _day(row["date"]) < segment_first
                        ]
                        derivation_rows = sorted(
                            [*context, *rows],
                            key=lambda row: (str(row["portid"]), _day(row["date"])),
                        )
                        _derive(derivation_rows, kind)

                    exported_by_year = defaultdict(list)
                    for row in rows:
                        pid = str(row["portid"])
                        day = _day(row["date"])
                        begin, end = wanted[pid]
                        if begin is None or end is None or not begin <= day <= end:
                            continue
                        target = stats[pid]
                        index = (day - begin).days
                        if target["present"][index]:
                            raise DownloadError(f"分块查询间存在重复节点日期：{pid}/{day}")
                        target["present"][index] = 1
                        target["records"] += 1
                        target["null_metric_cells"] += sum(row.get(metric) is None for metric in metrics)
                        target["first"] = day if target["first"] is None else min(target["first"], day)
                        target["last"] = day if target["last"] is None else max(target["last"], day)
                        node_row = selected_by_id[pid]
                        exported_by_year[day.year].append({
                            **row, "region": node_row["region"], "lat": node_row["lat"],
                            "lon": node_row["lon"], "source_url": pw.SOURCE,
                        })

                    if derived:
                        tail_start = segment_last - timedelta(days=28)
                        for pid in chunk:
                            candidates = [
                                row for row in derivation_rows
                                if str(row["portid"]) == pid
                                and tail_start <= _day(row["date"]) <= segment_last
                            ]
                            tails[pid] = candidates

                    for year, export_rows in sorted(exported_by_year.items()):
                        export_rows.sort(key=lambda row: (str(row["portid"]), str(row["date"])))
                        member = f"{kind}_daily_{year}_part{chunk_number:03d}.csv"
                        _write_csv_member(archive, member, export_rows, columns(kind, derived))
                        wrote_dataset = True
                    chunk_number += 1

            for node_row in selected:
                pid = str(node_row["portid"])
                target = stats[pid]
                begin, end = wanted[pid]
                limit = limits.get(pid)
                if mode == "latest" and limit and target["records"] != 1:
                    raise DownloadError(f"{pid} 的最新可用日未取到唯一记录")
                if mode == "all" and limit and target["records"] != limit["records"]:
                    raise DownloadError(f"{pid} 全历史记录数与节点统计不一致，请勾选重新读取源站后重试")
                present_count = target["records"]
                coverage.append({
                    "node_kind": kind, "portid": pid, "node_name": node_row["node_name"],
                    "status": "no_independent_statistics" if not node_row["statistics_available"] else
                             "no_records_in_requested_window" if not present_count else
                             "gaps" if present_count < target["expected"] else "complete",
                    "first_available_utc": limit["first"] if limit else None,
                    "latest_available_utc": limit["last"] if limit else None,
                    "export_first_utc": target["first"].isoformat() if target["first"] else None,
                    "export_last_utc": target["last"].isoformat() if target["last"] else None,
                    "records": present_count, "requested_days": target["expected"],
                    "missing_days": target["expected"] - present_count,
                    "missing_intervals_utc": _missing_intervals_bitmap(
                        begin, end, target["present"]
                    ),
                    "null_metric_cells": target["null_metric_cells"],
                })
            row_counts[kind] = sum(stats[pid]["records"] for pid in stats)
            if not wrote_dataset:
                _write_csv_member(archive, f"{kind}_daily.csv", [], columns(kind, derived))

        dictionary_rows = dictionary(derived)
        manifest = {
            "schema_version": VERSION, "source": "IMF PortWatch / UN Global Platform",
            "source_url": pw.SOURCE, "generated_at_utc": _now(), "timezone": "UTC",
            "mode": mode, "requested_start_utc": first.isoformat() if first else None,
            "requested_end_utc": last.isoformat() if last else None, "derived": derived,
            "node_count": sum(len(value[0]) for value in prepared.values()),
            "row_counts": row_counts, "fetch_complete": True,
            "coverage": coverage, "queries": metadata,
            "notes": ["最新指每节点自身最新可用日，保留实际日期及真实零值。",
                      "日表只含源站真实记录。缺失日不补零，见coverage.csv。",
                      "港口进出口是AIS推算货量（公吨）；要道capacity为通行运力（载重吨），不能当作实际贸易货量。",
                      "历史缓存最长24小时；queries记录实际抓取时间。上游数据会修订。",
                      "无独立统计的补充港口保留在节点目录。",
                      "7/30日均值须有连续窗口内全部有效数值，否则留空；前置日期只用于计算，不进入导出日表。",
                      "大范围结果按节点和日期分块写入ZIP，避免将整份历史表保留在Streamlit会话内存。",
                      "原始数字未转换成万吨；CSV以UTF-8 BOM保存，文本公式前缀加单引号避免误执行。",
                      "使用与再分发遵守IMF PortWatch来源条款；此下载不另授予数据许可。"],
        }
        archive.writestr("nodes.csv", csv_bytes(catalog, CATALOG_FIELDS))
        archive.writestr("coverage.csv", csv_bytes(coverage, COVERAGE_FIELDS))
        archive.writestr("data_dictionary.csv", csv_bytes(
            dictionary_rows, ["dataset", "field", "meaning", "unit", "missing"]
        ))
        archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2))
        archive.writestr("README.txt", "IMF PortWatch 日度下载\n"
                         + "\n".join(manifest["notes"])
                         + "\n来源：" + pw.SOURCE
                         + "\n生成UTC：" + manifest["generated_at_utc"])

    return {
        "datasets": {}, "artifact_path": str(artifact_path),
        "catalog": catalog, "coverage": coverage, "manifest": manifest,
        "dictionary": dictionary_rows,
    }


def collect(nodes: list[dict], mode: str, first: date | None = None, last: date | None = None,
            derived: bool = False, force: bool = False, progress=None) -> dict:
    if mode not in ("latest", "history", "all"):
        raise DownloadError("时间模式无效")
    if not nodes or any(n.get("node_kind") not in KINDS for n in nodes) or len(
        {(n['node_kind'], n['portid']) for n in nodes}) != len(nodes):
        raise DownloadError("请选择不重复的节点")
    if mode == "history" and (first is None or last is None or first > last or first < FIRST_DAY
                              or last > datetime.now(timezone.utc).date()):
        raise DownloadError("历史日期须在2019-01-01至今日UTC之间，开始日不得晚于结束日")
    prepared, estimated_rows = {}, 0
    for kind in KINDS:
        selected = [n for n in nodes if n["node_kind"] == kind]
        if not selected:
            continue
        ids = tuple(sorted(n["portid"] for n in selected if n["statistics_available"]))
        if progress:
            progress(f"查询{KINDS[kind]}各节点可用日期…")
        limits, queries = bounds(kind, ids)
        prepared[kind] = (selected, limits, queries)
        for node_row in selected:
            limit = limits.get(node_row["portid"])
            if not limit:
                continue
            if mode == "all":
                estimated_rows += limit["records"]
            elif mode == "latest":
                estimated_rows += 1
            else:
                begin, end = first, last
                estimated_rows += min(limit["records"], (end - begin).days + 1)

    if estimated_rows > INLINE_EXPORT_ROW_LIMIT:
        return _collect_streaming(prepared, mode, first, last, derived, force, progress)

    datasets, metadata, coverage, catalog = {}, [], [], []
    for kind in KINDS:
        if kind not in prepared:
            continue
        selected, limits, queries = prepared[kind]
        metadata.extend(queries)
        periods = defaultdict(list)
        wanted = {}
        for n in selected:
            limit = limits.get(n["portid"])
            entry = {**n, "first_available_utc": limit["first"] if limit else None,
                     "latest_available_utc": limit["last"] if limit else None}
            catalog.append(entry)
            if mode == "latest":
                begin = end = _day(limit["last"]) if limit else None
            elif mode == "all":
                begin, end = (_day(limit["first"]), _day(limit["last"])) if limit else (None, None)
            else:
                begin, end = first, last
            wanted[n["portid"]] = (begin, end)
            if not limit:
                continue
            fetch_first = max(_day(limit["first"]), begin - timedelta(days=29) if derived else begin)
            fetch_last = min(_day(limit["last"]), end)
            if fetch_first <= fetch_last:
                # Calendar years bound query size and make repeat downloads reusable.
                for a, b in _years(fetch_first, fetch_last):
                    for part_first, part_last, part_force in _refresh_segments(
                        a, b, force or mode == "latest"
                    ):
                        periods[(part_first, part_last, part_force)].append(n["portid"])
        reports, all_keys = [], set()
        for (begin, end, period_force), pids in sorted(periods.items()):
            for offset in range(0, len(pids), BATCH_SIZE):
                chunk = tuple(sorted(pids[offset:offset + BATCH_SIZE]))
                rows, query = fetch_window(kind, chunk, begin, end,
                    force=period_force,
                    progress=progress)
                metadata.append(query)
                for row in rows:
                    key = (row["portid"], row["date"])
                    if key in all_keys:
                        raise DownloadError("分块查询间存在重复节点日期")
                    all_keys.add(key)
                    reports.append(row)
        if derived:
            _derive(reports, kind)
        selected_by_id = {n["portid"]: n for n in selected}
        exported = []
        for row in reports:
            begin, end = wanted[row["portid"]]
            if begin <= _day(row["date"]) <= end:
                n = selected_by_id[row["portid"]]
                exported.append({**row, "region": n["region"], "lat": n["lat"], "lon": n["lon"],
                                 "source_url": pw.SOURCE})
        exported.sort(key=lambda r: (r["portid"], r["date"]))
        by_id = defaultdict(list)
        for row in exported:
            by_id[row["portid"]].append(row)
        metrics = PORT_METRICS if kind == "ports" else CHOKE_METRICS
        for n in selected:
            pid = n["portid"]
            rows, limit = by_id[pid], limits.get(pid)
            begin, end = wanted[pid]
            expected = (end - begin).days + 1 if begin and end else 0
            dates = {_day(r["date"]) for r in rows}
            if mode == "latest" and limit and len(rows) != 1:
                raise DownloadError(f"{pid} 的最新可用日未取到唯一记录")
            if mode == "all" and limit and len(rows) != limit["records"]:
                raise DownloadError(f"{pid} 全历史记录数与节点统计不一致，请勾选重新读取源站后重试")
            coverage.append({"node_kind": kind, "portid": pid, "node_name": n["node_name"],
                "status": "no_independent_statistics" if not n["statistics_available"] else
                          "no_records_in_requested_window" if not rows else
                          "gaps" if len(rows) < expected else "complete",
                "first_available_utc": limit["first"] if limit else None,
                "latest_available_utc": limit["last"] if limit else None,
                "export_first_utc": min(dates).isoformat() if dates else None,
                "export_last_utc": max(dates).isoformat() if dates else None,
                "records": len(rows), "requested_days": expected, "missing_days": expected - len(dates),
                "missing_intervals_utc": _missing_intervals(begin, end, dates) if begin and end else "",
                "null_metric_cells": sum(r.get(m) is None for r in rows for m in metrics)})
        datasets[kind] = exported
    manifest = {"schema_version": VERSION, "source": "IMF PortWatch / UN Global Platform",
                "source_url": pw.SOURCE, "generated_at_utc": _now(), "timezone": "UTC",
                "mode": mode, "requested_start_utc": first.isoformat() if first else None,
                "requested_end_utc": last.isoformat() if last else None, "derived": derived,
                "node_count": len(nodes), "row_counts": {k: len(v) for k, v in datasets.items()},
                "fetch_complete": True, "coverage": coverage, "queries": metadata,
                "notes": ["最新指每节点自身最新可用日，保留实际日期及真实零值。",
                          "日表只含源站真实记录。缺失日不补零，见coverage.csv。",
                          "港口进出口是AIS推算货量（公吨）；要道capacity为通行运力（载重吨），不能当作实际贸易货量。",
                          "历史缓存最长24小时；queries记录实际抓取时间。上游数据会修订。",
                          "无独立统计的补充港口保留在节点目录。",
                          "7/30日均值须有连续窗口内全部有效数值，否则留空；前置日期只用于计算，不进入导出日表。",
                          "原始数字未转换成万吨；CSV以UTF-8 BOM保存，文本公式前缀加单引号避免误执行。",
                          "使用与再分发遵守IMF PortWatch来源条款；此下载不另授予数据许可。"]}
    return {"datasets": datasets, "catalog": catalog, "coverage": coverage,
            "manifest": manifest, "dictionary": dictionary(derived)}


def columns(kind: str, derived: bool = False) -> list[str]:
    metrics = PORT_METRICS if kind == "ports" else CHOKE_METRICS
    extra = []
    if derived:
        mean_metrics = ["portcalls", "import", "export"] if kind == "ports" else ["n_total", "capacity"]
        for window in (7, 30):
            extra += [f"observed_days_{window}d"] + [f"avg_{metric}_{window}d" for metric in mean_metrics]
    return BASE_FIELDS + metrics + extra + ["source_url"]


def dictionary(derived: bool) -> list[dict]:
    entries = []
    labels = {"date": ("观测日", "UTC日期"), "portid": ("稳定节点编号", "文本"),
              "portname": ("源站节点名称", "文本"), "country": ("源站国家", "文本"),
              "ISO3": ("国家三位代码", "文本"), "region": ("项目监测水域", "文本"),
              "lat": ("节点纬度", "度"), "lon": ("节点经度", "度"), "source_url": ("数据来源", "URL")}
    for kind in KINDS:
        for field in columns(kind, derived):
            if field in labels:
                label, unit = labels[field]
            elif field.startswith("observed_days"):
                label, unit = "对应连续窗口内有源记录的天数", "天"
            else:
                metric = field.removeprefix("avg_")
                window = ""
                if field.startswith("avg_"):
                    metric, window = metric.rsplit("_", 1)
                prefix = next((p for p in ("portcalls", "import", "export", "capacity", "n_total", "n")
                               if metric == p or metric.startswith(p + "_")), "")
                names = {"portcalls": "有效进港", "import": "估算进口货量", "export": "估算出口货量",
                         "capacity": "通行运力", "n_total": "总通行", "n": "通行"}
                ship = metric[len(prefix):].lstrip("_")
                label = names.get(prefix, metric) + ("（" + TYPES[ship] + "）" if ship in TYPES else "")
                unit = "艘次" if prefix in ("portcalls", "n", "n_total") else "载重吨" if prefix == "capacity" else "公吨"
                if window:
                    label += f"，{window[:-1]}日均值"
                    unit += "/日"
            entries.append({"dataset": kind, "field": field, "meaning": label, "unit": unit,
                            "missing": "空值为未报告或窗口不完整；0为源站报告零值"})
    return entries


def zip_bytes(result: dict) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for kind, rows in result["datasets"].items():
            years = sorted({r["date"][:4] for r in rows})
            if len(years) > 1:
                for year in years:
                    archive.writestr(f"{kind}_daily_{year}.csv", csv_bytes(
                        [r for r in rows if r["date"].startswith(year)], columns(kind, result["manifest"]["derived"])))
            else:
                archive.writestr(f"{kind}_daily.csv", csv_bytes(rows, columns(kind, result["manifest"]["derived"])))
        archive.writestr("nodes.csv", csv_bytes(result["catalog"], CATALOG_FIELDS))
        archive.writestr("coverage.csv", csv_bytes(result["coverage"], COVERAGE_FIELDS))
        archive.writestr("data_dictionary.csv", csv_bytes(result["dictionary"], ["dataset", "field", "meaning", "unit", "missing"]))
        archive.writestr("manifest.json", json.dumps(result["manifest"], ensure_ascii=False, indent=2))
        archive.writestr("README.txt", "IMF PortWatch 日度下载\n" + "\n".join(result["manifest"]["notes"]) +
                         "\n来源：" + pw.SOURCE + "\n生成UTC：" + result["manifest"]["generated_at_utc"])
    return output.getvalue()


def xlsx_bytes(result: dict) -> bytes:
    """Build typed, filterable Excel sheets with the maintained openpyxl package."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    datasets = result["datasets"]
    records = sum(map(len, datasets.values()))
    if records > XLSX_ROW_LIMIT:
        raise DownloadError(
            f"Excel适合中小范围下载（最多{XLSX_ROW_LIMIT:,}条）；请用ZIP下载全量数据")

    manifest = result["manifest"]
    sheets = [
        ("港口日表" if kind == "ports" else "要道日表",
         columns(kind, manifest["derived"]), rows)
        for kind, rows in datasets.items()
    ]
    about = [
        {"项目": "来源", "内容": manifest["source"]},
        {"项目": "来源网址", "内容": manifest["source_url"]},
        {"项目": "生成时间（UTC）", "内容": manifest["generated_at_utc"]},
        {"项目": "时区", "内容": manifest["timezone"]},
        {"项目": "时间模式", "内容": manifest["mode"]},
        {"项目": "请求开始日（UTC）", "内容": manifest["requested_start_utc"]},
        {"项目": "请求结束日（UTC）", "内容": manifest["requested_end_utc"]},
        {"项目": "节点数", "内容": manifest["node_count"]},
        {"项目": "来源取数校验", "内容": "通过" if manifest["fetch_complete"] else "未通过"},
    ]
    about.extend({"项目": f"说明 {index}", "内容": note}
                 for index, note in enumerate(manifest["notes"], 1))
    queries = [
        {field: (json.dumps(query.get(field), ensure_ascii=False)
                 if isinstance(query.get(field), (dict, list)) else query.get(field))
         for field in ("operation", "endpoint", "where", "retrieved_at_utc",
                       "source_count", "exported_query_rows", "count_verified",
                       "cache_hit", "out_fields")}
        for query in manifest["queries"]
    ]
    sheets.extend([
        ("下载说明", ["项目", "内容"], about),
        ("节点目录", CATALOG_FIELDS, result["catalog"]),
        ("覆盖情况", COVERAGE_FIELDS, result["coverage"]),
        ("数据字典", ["dataset", "field", "meaning", "unit", "missing"], result["dictionary"]),
        ("来源查询", ["operation", "endpoint", "where", "retrieved_at_utc", "source_count",
                    "exported_query_rows", "count_verified", "cache_hit", "out_fields"], queries),
    ])

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="29435C")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet_index, (name, fields, rows) in enumerate(sheets, 1):
        worksheet = workbook.create_sheet(name)
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{max(1, len(rows) + 1)}"
        worksheet.append(fields)
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row in rows:
            values = []
            for field in fields:
                value = row.get(field)
                if isinstance(value, str):
                    # Keep untrusted text literal; openpyxl otherwise treats an
                    # initial '=' as an executable formula.
                    value = "".join(ch for ch in value if (
                        ch in "\t\n\r" or 0x20 <= ord(ch) <= 0xD7FF
                        or 0xE000 <= ord(ch) <= 0xFFFD
                        or 0x10000 <= ord(ch) <= 0x10FFFF))
                values.append(value)
            worksheet.append(values)
        for col_index, field in enumerate(fields, 1):
            column = get_column_letter(col_index)
            sample = [worksheet.cell(row, col_index).value
                      for row in range(1, min(worksheet.max_row, 100) + 1)]
            width = min(38, max(12, max((len(str(value)) for value in sample
                                        if value is not None), default=10) + 2))
            worksheet.column_dimensions[column].width = width
            for row_index in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row_index, col_index)
                if field == "date" and isinstance(cell.value, str):
                    try:
                        cell.value = date.fromisoformat(cell.value)
                        cell.number_format = "yyyy-mm-dd"
                    except ValueError:
                        pass
                elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    cell.number_format = "#,##0.###"
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        if rows:
            table = openpyxl.worksheet.table.Table(
                displayName=f"Data{sheet_index}",
                ref=f"A1:{get_column_letter(len(fields))}{len(rows) + 1}")
            table.tableStyleInfo = openpyxl.worksheet.table.TableStyleInfo(
                name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
            worksheet.add_table(table)

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
